from __future__ import annotations

import json
import time as time_module
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

import fitdecode
from openpyxl import load_workbook

from .config import Settings
from .db import connect, migrate, transaction
from .util import (
    atomic_write_json,
    ensure_aware,
    iso_utc,
    json_dumps,
    relative_to_root,
    sha256_file,
    stable_hash,
    write_heartbeat,
)


@dataclass
class IngestResult:
    imported_files: int = 0
    skipped_files: int = 0
    failed_files: int = 0
    health_records: int = 0
    activities: int = 0
    sensor_samples: int = 0

    def as_dict(self) -> dict[str, int]:
        return self.__dict__.copy()


def _begin_file(connection, settings: Settings, path: Path, parser_kind: str) -> tuple[int | None, str]:
    digest = sha256_file(path)
    parser_version = int(settings.values["ingest"].get("parser_version", 1))
    existing = connection.execute(
        "SELECT id, status FROM ingest_files WHERE sha256=? AND parser_kind=? AND parser_version=?",
        (digest, parser_kind, parser_version),
    ).fetchone()
    if existing and existing["status"] == "complete":
        return None, digest
    stat = path.stat()
    relative_path = relative_to_root(path, settings.root)
    now = iso_utc()
    if existing:
        connection.execute(
            """UPDATE ingest_files
               SET relative_path=?, size_bytes=?, mtime_ns=?, status='processing',
                   attempts=attempts+1, error_text=NULL, completed_at_utc=NULL
               WHERE id=?""",
            (relative_path, stat.st_size, stat.st_mtime_ns, existing["id"]),
        )
        return int(existing["id"]), digest
    cursor = connection.execute(
        """INSERT INTO ingest_files(
               relative_path, sha256, size_bytes, mtime_ns, parser_kind, parser_version,
               status, first_seen_at_utc
           ) VALUES(?,?,?,?,?,?,?,?)""",
        (relative_path, digest, stat.st_size, stat.st_mtime_ns, parser_kind, parser_version, "processing", now),
    )
    return int(cursor.lastrowid), digest


def _finish_file(connection, file_id: int) -> None:
    connection.execute(
        "UPDATE ingest_files SET status='complete', completed_at_utc=?, error_text=NULL WHERE id=?",
        (iso_utc(), file_id),
    )


def _fail_file(connection, file_id: int, error: Exception) -> None:
    connection.execute(
        "UPDATE ingest_files SET status='error', completed_at_utc=?, error_text=? WHERE id=?",
        (iso_utc(), f"{type(error).__name__}: {error}"[:4000], file_id),
    )


def _trim_header(value: Any) -> str:
    return " ".join(str(value).strip().split()) if value is not None else ""


def _cell_value(value: Any) -> Any:
    if value == "":
        return None
    if isinstance(value, float):
        # XLSX round-trips routinely perturb IEEE-754 tails without changing
        # the represented health value. Canonicalize only for hashes/raw JSON;
        # the typed metric column still stores the parser's numeric value.
        normalized = round(value, 8)
        return int(normalized) if normalized.is_integer() else normalized
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    return value


def _metric_value(value: Any) -> tuple[str, float | None, str | None]:
    if value is None or value == "":
        return "null", None, None
    if isinstance(value, bool):
        return "boolean", float(value), "true" if value else "false"
    if isinstance(value, timedelta):
        return "duration", value.total_seconds(), None
    if isinstance(value, datetime):
        return "datetime", None, value.isoformat()
    if isinstance(value, date):
        return "date", None, value.isoformat()
    if isinstance(value, time):
        return "datetime", None, value.isoformat()
    if isinstance(value, (int, float)):
        return "number", float(value), None
    return "text", None, str(value)


def _source_for_column(headers: list[str], values: list[Any], index: int) -> str | None:
    if index + 1 < len(headers) and headers[index + 1] == "Data Source":
        value = values[index + 1]
        return str(value) if value not in (None, "") else None
    trailing_sources = [values[i] for i, header in enumerate(headers) if header == "Data Source" and values[i] not in (None, "")]
    return str(trailing_sources[-1]) if len(trailing_sources) == 1 else None


def _excel_observed_times(sheet: str, row: dict[str, Any], timezone: ZoneInfo) -> tuple[str | None, str | None, str | None]:
    raw_date = row.get("Date")
    if isinstance(raw_date, datetime):
        base_date = raw_date.date()
    elif isinstance(raw_date, date):
        base_date = raw_date
    else:
        return None, None, None
    start_clock = row.get("Start") if sheet == "Sleep" else row.get("Time")
    end_clock = row.get("End") if sheet == "Sleep" else None
    if isinstance(start_clock, datetime):
        start_clock = start_clock.time()
    if not isinstance(start_clock, time):
        start_clock = time.min
    start_local = datetime.combine(base_date, start_clock, timezone)
    end_local: datetime | None = None
    if isinstance(end_clock, datetime):
        end_clock = end_clock.time()
    if isinstance(end_clock, time):
        end_local = datetime.combine(base_date, end_clock, timezone)
        if end_local <= start_local:
            end_local += timedelta(days=1)
    return iso_utc(start_local), iso_utc(end_local) if end_local else None, base_date.isoformat()


def _excel_natural_key(sheet: str, row: dict[str, Any]) -> str:
    identity_names = ("Date", "Time", "Start", "End", "Main", "Data Source")
    identity = {name: _cell_value(row.get(name)) for name in identity_names if row.get(name) not in (None, "")}
    if not identity:
        identity = {key: _cell_value(value) for key, value in row.items() if value not in (None, "")}
    return f"excel:{sheet}:{stable_hash(identity)}"


def ingest_excel(settings: Settings, connection, path: Path) -> IngestResult:
    result = IngestResult()
    with transaction(connection):
        file_id, _ = _begin_file(connection, settings, path, "health_xlsx")
    if file_id is None:
        result.skipped_files += 1
        return result
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        timezone = ZoneInfo(settings.timezone)
        catalog = settings.metric_catalog.get("metrics", {})
        with transaction(connection):
            for worksheet in workbook.worksheets:
                iterator = worksheet.iter_rows(values_only=True)
                headers_raw = next(iterator, None)
                if not headers_raw:
                    continue
                headers = [_trim_header(value) for value in headers_raw]
                last_used = max((i for i, header in enumerate(headers) if header), default=-1)
                headers = headers[: last_used + 1]
                for source_row, values_raw in enumerate(iterator, start=2):
                    values = list(values_raw[: len(headers)])
                    if not any(value not in (None, "") for value in values):
                        continue
                    row: dict[str, Any] = {}
                    for index, header in enumerate(headers):
                        if not header:
                            continue
                        key = header if header not in row else f"{header}#{index + 1}"
                        row[key] = values[index] if index < len(values) else None
                    record_key = _excel_natural_key(worksheet.title, row)
                    raw_payload = {key: _cell_value(value) for key, value in row.items()}
                    payload_hash = stable_hash(raw_payload)
                    current = connection.execute(
                        "SELECT id, revision, payload_hash FROM health_records WHERE record_key=? AND is_current=1",
                        (record_key,),
                    ).fetchone()
                    if current and current["payload_hash"] == payload_hash:
                        continue
                    revision = int(current["revision"] + 1) if current else 1
                    if current:
                        connection.execute("UPDATE health_records SET is_current=0 WHERE id=?", (current["id"],))
                    observed_start, observed_end, local_date = _excel_observed_times(worksheet.title, row, timezone)
                    source_names = [
                        str(values[i])
                        for i, header in enumerate(headers)
                        if header == "Data Source" and i < len(values) and values[i] not in (None, "")
                    ]
                    cursor = connection.execute(
                        """INSERT INTO health_records(
                               source_file_id, sheet_name, source_row, record_key, revision, is_current,
                               observed_start_utc, observed_end_utc, local_date, source_name,
                               raw_json, payload_hash, imported_at_utc
                           ) VALUES(?,?,?,?,?,1,?,?,?,?,?,?,?)""",
                        (
                            file_id,
                            worksheet.title,
                            source_row,
                            record_key,
                            revision,
                            observed_start,
                            observed_end,
                            local_date,
                            ", ".join(sorted(set(source_names))) or None,
                            json_dumps(raw_payload),
                            payload_hash,
                            iso_utc(),
                        ),
                    )
                    record_id = int(cursor.lastrowid)
                    result.health_records += 1
                    for index, header in enumerate(headers):
                        if not header or header in {"Date", "Time", "Start", "End", "Data Source"}:
                            continue
                        value = values[index] if index < len(values) else None
                        if value in (None, ""):
                            continue
                        definition = catalog.get(f"{worksheet.title}.{header}", {})
                        metric_key = str(definition.get("key") or f"excel.{worksheet.title}.{header}").lower().replace(" ", "_")
                        standard_unit = definition.get("unit")
                        value_type, value_number, value_text = _metric_value(value)
                        source_name = _source_for_column(headers, values, index)
                        connection.execute(
                            """INSERT INTO health_metrics(
                                   health_record_id, metric_key, display_name, value_type,
                                   value_number, value_text, raw_unit, standard_unit, source_name, raw_json
                               ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                            (
                                record_id,
                                metric_key,
                                header,
                                value_type,
                                value_number,
                                value_text,
                                standard_unit,
                                standard_unit,
                                source_name,
                                json_dumps({"value": _cell_value(value)}),
                            ),
                        )
                        if value_type == "number" and standard_unit is None and not definition.get("unitless", False):
                            existing_issue = connection.execute(
                                """SELECT 1 FROM data_quality_issues
                                   WHERE source_file_id=? AND entity_kind='health_metric'
                                     AND entity_key=? AND issue_code=?""",
                                (file_id, f"{worksheet.title}.{header}", "unconfirmed_metric_unit" if definition else "unknown_metric_unit"),
                            ).fetchone()
                            if not existing_issue:
                                connection.execute(
                                    """INSERT INTO data_quality_issues(
                                           source_file_id, entity_kind, entity_key, issue_code, severity, details_json, created_at_utc
                                       ) VALUES(?,?,?,?,?,?,?)""",
                                    (
                                        file_id,
                                        "health_metric",
                                        f"{worksheet.title}.{header}",
                                        "unconfirmed_metric_unit" if definition else "unknown_metric_unit",
                                        "warning",
                                        json_dumps({"sheet": worksheet.title, "column": header}),
                                        iso_utc(),
                                    ),
                                )
            _finish_file(connection, file_id)
        workbook.close()
        result.imported_files += 1
    except Exception as error:
        with transaction(connection):
            _fail_file(connection, file_id, error)
        result.failed_files += 1
        raise
    return result


def _fit_dict(frame: fitdecode.FitDataMessage) -> dict[str, tuple[Any, str | None]]:
    return {field.name: (field.value, field.units) for field in frame.fields}


def _fit_raw(fields: dict[str, tuple[Any, str | None]]) -> dict[str, Any]:
    return {key: {"value": _cell_value(value), "unit": unit} for key, (value, unit) in fields.items()}


def _uuid_hex(value: Any) -> str | None:
    if not isinstance(value, (tuple, list)) or not value:
        return None
    try:
        return bytes(int(part) for part in value).hex()
    except (TypeError, ValueError):
        return None


def _canonical_sport(sport: Any, sub_sport: Any) -> tuple[str, str | None]:
    raw_sport = str(sport or "unknown")
    raw_sub = str(sub_sport) if sub_sport is not None else None
    if raw_sport == "running":
        return "running", raw_sub
    if raw_sport in {"rock_climbing", "climbing"}:
        return "climbing", raw_sub
    if raw_sport == "training" and raw_sub == "strength_training":
        return "strength_training", raw_sub
    return raw_sport, raw_sub


def _fit_datetime(value: Any) -> datetime | None:
    if not isinstance(value, datetime):
        return None
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def _insert_activity_metric(connection, activity_id: int, segment_id: int | None, key: str, value: Any, unit: str | None) -> None:
    if value is None:
        return
    value_type, value_number, value_text = _metric_value(value)
    if isinstance(value, (tuple, list, dict)):
        value_type, value_number, value_text = "text", None, json_dumps(value)
    connection.execute(
        """INSERT OR REPLACE INTO activity_metrics(
               activity_id, segment_id, metric_key, value_type, value_number, value_text, unit, raw_json
           ) VALUES(?,?,?,?,?,?,?,?)""",
        (activity_id, segment_id, key, value_type, value_number, value_text, unit, json_dumps({"value": _cell_value(value)})),
    )


def ingest_fit(settings: Settings, connection, path: Path) -> IngestResult:
    result = IngestResult()
    with transaction(connection):
        file_id, _ = _begin_file(connection, settings, path, "fit")
    if file_id is None:
        result.skipped_files += 1
        return result
    try:
        sessions: list[dict[str, tuple[Any, str | None]]] = []
        laps: list[dict[str, tuple[Any, str | None]]] = []
        zones: list[dict[str, tuple[Any, str | None]]] = []
        records: list[dict[str, tuple[Any, str | None]]] = []
        device: dict[str, tuple[Any, str | None]] = {}
        sport_message: dict[str, tuple[Any, str | None]] = {}
        with fitdecode.FitReader(path) as reader:
            for frame in reader:
                if not isinstance(frame, fitdecode.FitDataMessage):
                    continue
                fields = _fit_dict(frame)
                if frame.name == "session":
                    sessions.append(fields)
                elif frame.name == "lap":
                    laps.append(fields)
                elif frame.name == "time_in_zone":
                    zones.append(fields)
                elif frame.name == "record":
                    records.append(fields)
                elif frame.name == "device_info" and not device:
                    device = fields
                elif frame.name == "sport" and not sport_message:
                    sport_message = fields
        if not sessions:
            raise ValueError("FIT file has no session message")
        timezone = ZoneInfo(settings.timezone)
        device_name = str(device.get("product_name", (None, None))[0] or "") or None
        source_application = str(device.get("descriptor", (None, None))[0] or "") or None
        inserted: list[tuple[int, datetime, datetime | None]] = []
        with transaction(connection):
            for session_index, fields in enumerate(sessions):
                start = _fit_datetime(fields.get("start_time", (None, None))[0])
                end = _fit_datetime(fields.get("timestamp", (None, None))[0])
                if start is None:
                    raise ValueError(f"FIT session {session_index} has no start_time")
                raw_sport = fields.get("sport", sport_message.get("sport", ("unknown", None)))[0]
                raw_sub = fields.get("sub_sport", sport_message.get("sub_sport", (None, None)))[0]
                sport_type, sport_subtype = _canonical_sport(raw_sport, raw_sub)
                duration = fields.get("total_timer_time", fields.get("total_elapsed_time", (None, None)))[0]
                duration_number = float(duration) if isinstance(duration, (int, float)) else None
                uuid_value = fields.get("SESSION UUID", (None, None))[0]
                session_uuid = _uuid_hex(uuid_value)
                fallback = stable_hash(
                    {
                        "device": device_name,
                        "start": start.isoformat(),
                        "sport": sport_type,
                        "duration": round(duration_number or 0, 3),
                    }
                )
                source_key = f"fit:uuid:{session_uuid}" if session_uuid else f"fit:fingerprint:{fallback}"
                existing = connection.execute("SELECT id FROM activities WHERE source_key=?", (source_key,)).fetchone()
                if existing:
                    inserted.append((int(existing["id"]), start, end))
                    continue
                start_local = start.astimezone(timezone)
                cursor = connection.execute(
                    """INSERT INTO activities(
                           source_file_id, source_key, fit_session_uuid, fallback_fingerprint,
                           sport_type, sport_subtype, start_time_utc, end_time_utc, start_time_local,
                           local_date, duration_seconds, source_application, device_name, raw_json, imported_at_utc
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        file_id,
                        source_key,
                        session_uuid,
                        fallback,
                        sport_type,
                        sport_subtype,
                        iso_utc(start),
                        iso_utc(end) if end else None,
                        start_local.isoformat(),
                        start_local.date().isoformat(),
                        duration_number,
                        source_application,
                        device_name,
                        json_dumps(_fit_raw(fields)),
                        iso_utc(),
                    ),
                )
                activity_id = int(cursor.lastrowid)
                inserted.append((activity_id, start, end))
                result.activities += 1
                for key, (value, unit) in fields.items():
                    if key not in {"start_time", "timestamp", "sport", "sub_sport", "SESSION UUID"}:
                        _insert_activity_metric(connection, activity_id, None, key, value, unit)

            for lap_index, fields in enumerate(laps):
                lap_start = _fit_datetime(fields.get("start_time", (None, None))[0])
                activity = _activity_for_time(inserted, lap_start)
                if not activity:
                    continue
                activity_id = activity[0]
                lap_end = _fit_datetime(fields.get("timestamp", (None, None))[0])
                duration = fields.get("total_timer_time", (None, None))[0]
                cursor = connection.execute(
                    """INSERT OR IGNORE INTO activity_segments(
                           activity_id, segment_type, segment_index, start_time_utc, end_time_utc, duration_seconds, raw_json
                       ) VALUES(?,?,?,?,?,?,?)""",
                    (
                        activity_id,
                        "lap",
                        lap_index,
                        iso_utc(lap_start) if lap_start else None,
                        iso_utc(lap_end) if lap_end else None,
                        float(duration) if isinstance(duration, (int, float)) else None,
                        json_dumps(_fit_raw(fields)),
                    ),
                )
                segment = connection.execute(
                    "SELECT id FROM activity_segments WHERE activity_id=? AND segment_type='lap' AND segment_index=?",
                    (activity_id, lap_index),
                ).fetchone()
                segment_id = int(segment["id"])
                for key, (value, unit) in fields.items():
                    if key not in {"start_time", "timestamp"}:
                        _insert_activity_metric(connection, activity_id, segment_id, key, value, unit)

            for zone_index, fields in enumerate(zones):
                timestamp = _fit_datetime(fields.get("timestamp", (None, None))[0])
                activity = _activity_for_time(inserted, timestamp)
                if not activity:
                    continue
                activity_id = activity[0]
                connection.execute(
                    """INSERT OR IGNORE INTO activity_segments(
                           activity_id, segment_type, segment_index, start_time_utc, end_time_utc, duration_seconds, raw_json
                       ) VALUES(?,?,?,?,?,?,?)""",
                    (activity_id, "heart_rate_zones", zone_index, None, iso_utc(timestamp) if timestamp else None, None, json_dumps(_fit_raw(fields))),
                )

            for sample_index, fields in enumerate(records):
                timestamp = _fit_datetime(fields.get("timestamp", (None, None))[0])
                activity = _activity_for_time(inserted, timestamp)
                if not activity or timestamp is None:
                    continue
                activity_id = activity[0]
                for key, (value, unit) in fields.items():
                    if key == "timestamp" or value is None:
                        continue
                    if isinstance(value, bool):
                        value_type, number, text = "boolean", float(value), "true" if value else "false"
                    elif isinstance(value, (int, float)):
                        value_type, number, text = "number", float(value), None
                    else:
                        value_type, number, text = "text", None, json_dumps(value) if isinstance(value, (tuple, list, dict)) else str(value)
                    connection.execute(
                        """INSERT OR IGNORE INTO sensor_samples(
                               activity_id, timestamp_utc, sample_index, metric_key, value_type,
                               value_number, value_text, unit, raw_json
                           ) VALUES(?,?,?,?,?,?,?,?,?)""",
                        (activity_id, iso_utc(timestamp), sample_index, key, value_type, number, text, unit, json_dumps({"value": _cell_value(value)})),
                    )
                    result.sensor_samples += 1
            _finish_file(connection, file_id)
        result.imported_files += 1
    except Exception as error:
        with transaction(connection):
            _fail_file(connection, file_id, error)
        result.failed_files += 1
        raise
    return result


def _activity_for_time(
    activities: list[tuple[int, datetime, datetime | None]], timestamp: datetime | None
) -> tuple[int, datetime, datetime | None] | None:
    if not activities:
        return None
    if timestamp is None or len(activities) == 1:
        return activities[0]
    for activity in activities:
        _, start, end = activity
        if timestamp >= start and (end is None or timestamp <= end + timedelta(minutes=1)):
            return activity
    return min(activities, key=lambda item: abs((timestamp - item[1]).total_seconds()))


def ingest_once(settings: Settings, *, only_paths: set[Path] | None = None) -> IngestResult:
    combined = IngestResult()
    connection = connect(settings.database_path, busy_timeout_ms=int(settings.values["sqlite"].get("busy_timeout_ms", 10_000)))
    migrate(connection)
    paths: list[tuple[Path, str]] = [(settings.path("health_workbook"), "excel")]
    fit_dir = settings.path("fit_directory")
    if fit_dir.exists():
        paths.extend((path, "fit") for path in sorted(fit_dir.glob("*.fit")))
    for path, kind in paths:
        if not path.is_file():
            continue
        if only_paths is not None and path.resolve() not in {item.resolve() for item in only_paths}:
            continue
        try:
            current = ingest_excel(settings, connection, path) if kind == "excel" else ingest_fit(settings, connection, path)
        except Exception:
            combined.failed_files += 1
            continue
        for field in combined.__dict__:
            setattr(combined, field, getattr(combined, field) + getattr(current, field))
    connection.close()
    atomic_write_json(
        settings.path("state_directory") / "ingest_check.json",
        {"checked_at_utc": iso_utc(), "result": combined.as_dict()},
    )
    return combined


def _stable_files(paths: Iterable[Path], stability_seconds: int) -> list[Path]:
    before = {path: (path.stat().st_size, path.stat().st_mtime_ns) for path in paths if path.is_file()}
    if stability_seconds > 0:
        time_module.sleep(stability_seconds)
    return [
        path
        for path, signature in before.items()
        if path.is_file() and (path.stat().st_size, path.stat().st_mtime_ns) == signature
    ]


def ingest_daemon(settings: Settings) -> None:
    heartbeat = settings.path("state_directory") / "heartbeats" / "ingest.json"
    poll_seconds = int(settings.values["ingest"].get("poll_seconds", 15))
    stability_seconds = int(settings.values["ingest"].get("stability_seconds", 10))
    while True:
        workbook = settings.path("health_workbook")
        fit_dir = settings.path("fit_directory")
        candidates = ([workbook] if workbook.exists() else []) + (sorted(fit_dir.glob("*.fit")) if fit_dir.exists() else [])
        stable = _stable_files(candidates, stability_seconds)
        try:
            result = ingest_once(settings, only_paths=set(stable)) if stable else IngestResult()
            write_heartbeat(heartbeat, status="ok", details=result.as_dict())
        except Exception as error:
            write_heartbeat(heartbeat, status="error", details={"error": f"{type(error).__name__}: {error}"})
        time_module.sleep(max(1, poll_seconds))
