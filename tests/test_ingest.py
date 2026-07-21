from __future__ import annotations

import copy
import shutil
from dataclasses import replace

from openpyxl import load_workbook

from trainlab.db import connect, table_counts
from trainlab.ingest import ingest_fit, ingest_once

from .conftest import make_settings


def test_golden_sample_counts_units_and_time(settings):
    connection = connect(settings.database_path)
    counts = table_counts(connection)
    assert counts["ingest_files"] == 5
    assert counts["health_records"] == 1011
    assert counts["activities"] == 4
    assert counts["sensor_samples"] == 42762
    run = connection.execute("SELECT * FROM activities WHERE sport_type='running'").fetchone()
    assert run["local_date"] == "2026-04-15"
    assert run["start_time_local"].startswith("2026-04-15T06:53:33+08:00")
    distance = connection.execute(
        "SELECT value_number, unit FROM activity_metrics WHERE activity_id=? AND metric_key='total_distance' AND segment_id IS NULL",
        (run["id"],),
    ).fetchone()
    assert distance["value_number"] == 3913.62
    assert distance["unit"] == "m"
    assert connection.execute("SELECT COUNT(*) FROM activity_segments WHERE activity_id=? AND segment_type='lap'", (run["id"],)).fetchone()[0] == 4
    assert connection.execute("SELECT COUNT(*) FROM sensor_samples WHERE activity_id=? AND metric_key='cadence'", (run["id"],)).fetchone()[0] > 0
    connection.close()


def test_all_eight_health_workbook_sheets_have_golden_counts(settings):
    connection = connect(settings.database_path)
    counts = {
        row["sheet_name"]: row["count"]
        for row in connection.execute(
            "SELECT sheet_name, COUNT(*) AS count FROM health_records WHERE is_current=1 GROUP BY sheet_name"
        )
    }
    assert counts == {
        "Daily Metrics": 360,
        "Weight": 16,
        "Sleep": 107,
        "Nutrition": 177,
        "Mindfulness": 351,
    }
    assert {"Blood Sugar", "Blood Pressure", "Womens Health"}.isdisjoint(counts)
    connection.close()


def test_duplicate_scan_and_renamed_fit_are_idempotent(settings, tmp_path):
    before = connect(settings.database_path)
    counts_before = table_counts(before)
    before.close()
    result = ingest_once(settings)
    assert result.skipped_files == 5
    renamed = tmp_path / "renamed.fit"
    original = next(settings.path("fit_directory").glob("*.fit"))
    shutil.copy2(original, renamed)
    connection = connect(settings.database_path)
    renamed_result = ingest_fit(settings, connection, renamed)
    assert renamed_result.skipped_files == 1
    assert table_counts(connection) == counts_before
    connection.close()


def test_excel_change_creates_one_revision(golden_settings, tmp_path):
    workbook_path = tmp_path / "Health.xlsx"
    shutil.copy2(golden_settings.path("health_workbook"), workbook_path)
    fit_dir = tmp_path / "empty-fit"
    fit_dir.mkdir()
    current = make_settings(golden_settings, tmp_path / "runtime", fit_directory=fit_dir, workbook=workbook_path)
    first = ingest_once(current)
    assert first.health_records == 1011
    workbook = load_workbook(workbook_path)
    sheet = workbook["Daily Metrics"]
    old_value = sheet.cell(2, 6).value
    sheet.cell(2, 6).value = float(old_value) + 1
    workbook.save(workbook_path)
    second = ingest_once(current)
    assert second.health_records == 1
    connection = connect(current.database_path)
    assert connection.execute("SELECT COUNT(*) FROM health_records WHERE revision=2").fetchone()[0] == 1
    assert connection.execute("SELECT COUNT(*) FROM health_records WHERE is_current=1").fetchone()[0] == 1011
    connection.close()


def test_corrupt_fit_is_audited_and_retryable(base_settings, tmp_path):
    fit_dir = tmp_path / "HealthFit"
    fit_dir.mkdir()
    broken = fit_dir / "broken.fit"
    broken.write_bytes(b"partial-fit")
    workbook = tmp_path / "missing.xlsx"
    current = make_settings(base_settings, tmp_path / "runtime", fit_directory=fit_dir, workbook=workbook)
    first = ingest_once(current)
    second = ingest_once(current)
    assert first.failed_files == 1
    assert second.failed_files == 1
    connection = connect(current.database_path)
    row = connection.execute("SELECT status, attempts, error_text FROM ingest_files").fetchone()
    assert row["status"] == "error"
    assert row["attempts"] == 2
    assert row["error_text"]
    connection.close()


def test_unknown_numeric_metric_produces_one_quality_warning(base_settings, tmp_path):
    workbook_path = tmp_path / "Health.xlsx"
    shutil.copy2(base_settings.path("health_workbook"), workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Daily Metrics"]
    sheet.cell(1, 10).value = "Future Sensor"
    sheet.cell(2, 10).value = 12.3
    sheet.cell(3, 10).value = 13.4
    workbook.save(workbook_path)
    fit_dir = tmp_path / "empty-fit"
    fit_dir.mkdir()
    current = make_settings(base_settings, tmp_path / "runtime", fit_directory=fit_dir, workbook=workbook_path)
    ingest_once(current)
    connection = connect(current.database_path)
    issues = connection.execute(
        "SELECT entity_key, issue_code FROM data_quality_issues WHERE entity_key='Daily Metrics.Future Sensor'"
    ).fetchall()
    assert len(issues) == 1
    assert issues[0]["issue_code"] == "unknown_metric_unit"
    connection.close()
